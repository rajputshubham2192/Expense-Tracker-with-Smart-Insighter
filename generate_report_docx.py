import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def create_report_docx(output_filename="PROJECT_OVERALL_REPORT.docx"):
    doc = Document()

    # Set page margins (0.75 inch)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # slate-700
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("PROJECT OVERALL REPORT")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # dark slate

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Expense Tracker with Smart Insighter — AI/ML Financial Intelligence & Analytics System")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # royal blue

    # Metadata Table
    t_meta = doc.add_table(rows=2, cols=4)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_meta, color="E2E8F0", sz="4")
    
    meta_items = [
        ("Project Title:", "Expense Tracker with Smart Insighter", "Date:", "September 25, 2026"),
        ("Prepared By:", "Shubham Kumar Singh", "Architecture:", "FastAPI + React 18 + Scikit-Learn")
    ]

    col_widths = [Inches(1.2), Inches(2.3), Inches(1.1), Inches(2.4)]

    for row_idx, row_data in enumerate(meta_items):
        for col_idx in range(4):
            cell = t_meta.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "F8FAFC" if col_idx in [0, 2] else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(row_data[col_idx])
            r.font.size = Pt(9.5)
            if col_idx in [0, 2]:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            else:
                r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def add_section_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if level == 1:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif level == 2:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r_b = p.add_run(bold_prefix + " ")
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_t = p.add_run(text)
        r_t.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 1. Executive Summary
    add_section_heading("1. Executive Summary", 1)
    p = doc.add_paragraph(
        "Expense Tracker with Smart Insighter is an enterprise-grade, full-stack personal finance intelligence system "
        "designed to transform personal money management from passive, retroactive bookkeeping into an active, proactive financial advisory platform. "
        "Built with FastAPI (Python) on the backend and React 18 (Vite) on the frontend, the platform natively integrates machine learning models "
        "(Natural Language Processing, Isolation Forest Outlier Detection, Time-Series Velocity Regression, and K-Means Behavioral Clustering) "
        "within a sleek dark-mode glassmorphic interface."
    )

    # 2. Why This Project? (Problem Statement & Motivation)
    add_section_heading("2. Motivation & Problem Statement (The 'Why')", 1)
    doc.add_paragraph(
        "Most existing personal finance trackers operate as digital notebooks that record past actions without delivering contextual foresight. "
        "The primary industry pain points resolved by this platform include:"
    )
    add_bullet("High Manual Friction & Tracking Fatigue:", "Manual entry and category tagging for every purchase lead to user drop-off. Smart Insighter automates ingestion via NLP classification.")
    add_bullet("Absence of Predictive Foresight:", "Traditional retroactive charts show overspending after it occurs. Smart Insighter projects month-end balances 10–15 days in advance using velocity modeling.")
    add_bullet("Hidden Subscription Creep:", "Recurring services (streaming, gym, SaaS) often raise prices without clear warnings. Smart Insighter actively tracks recurring cycles and price inflation.")
    add_bullet("Unidentified Spending Anomalies:", "Spikes from accidental duplicate transactions or billing errors get lost in large spreadsheets. The Isolation Forest radar isolates statistical outliers in real time.")
    add_bullet("Lack of System-Wide Governance:", "Administrators lack high-level visibility into user engagement and gross platform volume. The platform features an Administrator Oversight Portal for centralized telemetry.")

    # 3. System Architecture & Tech Stack
    add_section_heading("3. System Architecture & Technical Stack (The 'How')", 1)
    doc.add_paragraph(
        "The platform utilizes a modern decoupled client-server architecture with high-performance asynchronous endpoints and dedicated machine learning service layers:"
    )

    # Tech Stack Table
    t_tech = doc.add_table(rows=1, cols=3)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tech, color="CBD5E1", sz="4")

    headers = ["Layer", "Technology", "Role & Engineering Highlights"]
    hdr_widths = [Inches(1.5), Inches(2.0), Inches(3.5)]

    # Header Row
    for i, title in enumerate(headers):
        cell = t_tech.cell(0, i)
        cell.width = hdr_widths[i]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tech_rows = [
        ("Frontend Client", "React 18, Vite, Lucide React", "Single Page Application (SPA) with modular hooks, dark-mode glassmorphism, and responsive flex/grid layouts."),
        ("Backend Server", "FastAPI, Python 3.10+, Uvicorn", "Asynchronous high-throughput REST API with automatic OpenAPI documentation and strict Pydantic schema validation."),
        ("Database & ORM", "SQLite / PostgreSQL, SQLAlchemy 2.0", "ACID-compliant relational persistence, indexed foreign key relations, and declarative ORM models."),
        ("Machine Learning", "Scikit-Learn, NumPy, Pandas", "TF-IDF text vectorization, supervised categorizer, Isolation Forest anomaly scanner, and velocity regression."),
        ("Authentication & Security", "JWT (python-jose), Passlib (Bcrypt)", "Stateless bearer token authentication with encrypted claims and salted password hashing (12+ rounds)."),
        ("Deployment & Hosting", "Vercel (Client), Cloud Server (API)", "Production-optimized build with client-side SPA routing rewrites (vercel.json) and flexible baseURL client setup.")
    ]

    for row_idx, r_data in enumerate(tech_rows):
        row = t_tech.add_row()
        for c_idx, text in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = hdr_widths[c_idx]
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.0)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Deep Dive: Machine Learning Intelligence Core
    add_section_heading("4. Machine Learning & Intelligence Modules", 1)
    
    add_section_heading("4.1 NLP Transaction Auto-Categorization (ml_categorizer.py)", 2)
    doc.add_paragraph(
        "Transforms merchant descriptions (e.g. 'Starbucks Coffee', 'Uber Ride', 'Netflix Premium') into numerical token matrices using "
        "TF-IDF (Term Frequency - Inverse Document Frequency) with n-gram range (1, 2). A supervised classifier maps the vectorized tokens to standardized "
        "categories and computes a confidence probability score. When users manually correct a category, the system logs feedback for continuous adaptive retraining."
    )

    add_section_heading("4.2 Time-Series Velocity Forecaster (ml_forecasting.py)", 2)
    doc.add_paragraph(
        "Analyzes daily cumulative spending trends to compute spending velocity and extrapolate month-end burn rate:\n"
        "• Daily Burn Velocity = Cumulative Spend / Current Day of Month\n"
        "• Projected Month-End Spend = Current Spend + (Daily Burn Velocity × Remaining Days in Month)\n"
        "This dynamic forecast is compared against active budget ceilings to alert the user with tiered warnings (Safe, Warning, Risk of Breach)."
    )

    add_section_heading("4.3 Isolation Forest Anomaly Detection (ml_anomaly.py)", 2)
    doc.add_paragraph(
        "Employs an unsupervised Isolation Forest tree ensemble algorithm alongside statistical Z-scores across spending amount, frequency, and category variances. "
        "Outlier data points isolate in fewer tree partitions and receive high anomaly scores, instantly alerting the user to abnormal spikes or potential billing discrepancies."
    )

    add_section_heading("4.4 Subscription Hike Radar (recommendation_engine.py)", 2)
    doc.add_paragraph(
        "Clusters transactions by merchant and analyzes billing intervals (28–32 days). It identifies active recurring subscriptions and detects price increases across billing cycles."
    )

    add_section_heading("4.5 Multi-Format CSV Bank Statement Ingestion (csv_parser.py)", 2)
    doc.add_paragraph(
        "Uses fuzzy header detection to automatically map diverse bank statement formats (Date, Amount, Description, Balance) and batch-processes imported records through the NLP auto-categorization pipeline."
    )

    # 5. Application Modules & Capabilities
    add_section_heading("5. Core Modules & User Capabilities", 1)
    add_bullet("User Dashboard & KPI Suite:", "Displays monthly inflow, outflow, net savings, financial health score (0–100), and interactive 60-day spending activity heatmap matrix.")
    add_bullet("Transaction Management:", "Full CRUD operations with multi-criteria filtering (date range, category, payment mode) and CSV statement importing.")
    add_bullet("Smart Budgets & Safeguards:", "Category-wise and overall monthly budget ceilings with visual progress indicators (Safe < 75%, Warning 75%–99%, Exceeded ≥ 100%).")
    add_bullet("AI Smart Insighter Hub:", "Dedicated analytics center presenting month-end spend forecasts, flagged transaction anomalies, subscription radars, and personalized savings advice.")
    add_bullet("Administrator Oversight Portal:", "Centralized platform telemetry displaying total registered users, active user directory, toggle controls, platform gross volume, and aggregated category distributions.")

    # 6. Security & Engineering Standards
    add_section_heading("6. Security, Resilience & Quality Standards", 1)
    add_bullet("Stateless JWT Authentication:", "Protected routes verify cryptographically signed tokens on every request via FastAPI dependency injection.")
    add_bullet("Bcrypt Password Hashing:", "Salted password hashing (12+ rounds) ensures credentials are secure against rainbow table attacks.")
    add_bullet("Role-Based Access Control (RBAC):", "Strict middleware enforcement separates regular user scopes from administrative system capabilities.")
    add_bullet("Production Ready Deployment:", "Configured with vercel.json rewrite routing for seamless SPA reloads and environment-based API base URLs.")

    # 7. Business Impact & Conclusion
    add_section_heading("7. Conclusion & Value Proposition", 1)
    doc.add_paragraph(
        "Expense Tracker with Smart Insighter successfully modernizes personal financial management. "
        "By fusing modern web architecture with machine learning algorithms, the platform replaces static tracking sheets with an automated, "
        "predictive, and secure financial intelligence hub that empowers users to eliminate overspending, monitor subscriptions, and achieve sustainable financial wellness."
    )

    doc.save(output_filename)
    print(f"Word Document successfully created: {output_filename}")

if __name__ == "__main__":
    create_report_docx()
